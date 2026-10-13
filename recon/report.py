"""Excel exception report writer."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .engine import ReconResult

NAVY = "1F2A44"
STATUS_FILL = {
    "Quantity break": "F8D7DA",
    "Missing internally": "F8D7DA",
    "Missing at custodian": "F8D7DA",
    "Cash break": "FCE5CD",
    "Price break": "FFF3CD",
    "Matched – review ID": "DDEBF7",
    "Matched": "E2F0D9",
}
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NUM_FORMATS = {
    "Internal Qty": "#,##0.####", "Custodian Qty": "#,##0.####", "Qty Diff": "+#,##0.####;-#,##0.####;0",
    "Internal Price": "#,##0.00", "Custodian Price": "#,##0.00", "Price Diff %": "+0.00;-0.00;0.00",
    "Internal MV": "#,##0.00", "Custodian MV": "#,##0.00", "MV Diff": "+#,##0.00;-#,##0.00;0.00",
}
WIDTHS = {"Status": 22, "Security": 28, "Suggested Cause": 70, "Match Method": 14, "ISIN": 15}


def _header(ws, row, cols):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _write_table(ws, df: pd.DataFrame, start_row: int = 1):
    cols = list(df.columns)
    _header(ws, start_row, cols)
    for r, rec in enumerate(df.itertuples(index=False), start_row + 1):
        status = rec[cols.index("Status")] if "Status" in cols else None
        for j, val in enumerate(rec, 1):
            if isinstance(val, float) and pd.isna(val):
                val = None
            cell = ws.cell(row=r, column=j, value=val)
            cell.border = BORDER
            col = cols[j - 1]
            if col in NUM_FORMATS:
                cell.number_format = NUM_FORMATS[col]
            if col == "Suggested Cause":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
            if status and col == "Status":
                cell.fill = PatternFill("solid", fgColor=STATUS_FILL.get(status, "FFFFFF"))
                cell.font = Font(bold=True)
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(c, max(12, len(c) + 2))
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(df)}"


def build_excel(result: ReconResult, internal_name: str, custodian_name: str, as_of: str | None = None) -> bytes:
    s = result.summary
    tol = s["tolerances"]
    wb = Workbook()

    # ---- Summary sheet -----------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Portfolio Reconciliation Report"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    meta = [
        ("Run at", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("As-of date", as_of or "—"),
        ("Internal file", internal_name),
        ("Custodian file", custodian_name),
        ("Tolerances", f"Qty ±{tol.qty_abs:g} sh · Price ±{tol.price_pct:g}% · Cash ±{tol.cash_abs:g} · Name match ≥{tol.fuzzy_score}%"),
    ]
    for k, (a, b) in enumerate(meta, 3):
        ws.cell(row=k, column=1, value=a).font = Font(bold=True)
        ws.cell(row=k, column=2, value=b)

    r = 3 + len(meta) + 1
    ws.cell(row=r, column=1, value="Result").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    _header(ws, r, ["Status", "Count"])
    for status, n in s["by_status"].items():
        r += 1
        c1 = ws.cell(row=r, column=1, value=status)
        c1.fill = PatternFill("solid", fgColor=STATUS_FILL.get(status, "FFFFFF"))
        c1.border = BORDER
        ws.cell(row=r, column=2, value=n).border = BORDER
    r += 1
    ws.cell(row=r, column=1, value="Total positions compared").font = Font(bold=True)
    ws.cell(row=r, column=2, value=len(result.detail)).font = Font(bold=True)

    r += 2
    ws.cell(row=r, column=1, value="Market value by currency").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    mv = s["mv_by_ccy"].reset_index()
    mv["Gross Break MV"] = mv["Currency"].map(s["abs_break_mv_by_ccy"]).fillna(0)
    _header(ws, r, list(mv.columns))
    for rec in mv.itertuples(index=False):
        r += 1
        for j, val in enumerate(rec, 1):
            cell = ws.cell(row=r, column=j, value=val)
            cell.border = BORDER
            if j > 1:
                cell.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 26
    for col in "BCDE":
        ws.column_dimensions[col].width = 18  # long meta text overflows into empty cells to the right

    # ---- Breaks & full detail ---------------------------------------------
    breaks = result.detail[result.detail["Status"] != "Matched"]
    _write_table(wb.create_sheet("Breaks"), breaks)
    _write_table(wb.create_sheet("All Positions"), result.detail)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
